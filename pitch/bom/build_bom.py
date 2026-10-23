#!/usr/bin/env python3
"""
Build a TheWebxite bill of materials as a matched .xlsx + .pdf pair.

    python3 build_bom.py --spec bom.json --outdir /mnt/user-data/outputs

The Excel is the editable master; the PDF is generated from it so the two can
never drift apart. See SKILL.md for the spec format.
"""

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NAVY = "1F3864"
LTBLUE = "EAF0F7"
YELLOW = "FFF2CC"
AMBER = "FCE4D6"
GREY = "F2F2F2"
GREEN = "E2EFDA"
DKGREEN = "375623"
DKAMBER = "833C0C"
GREYTEXT = "404040"

RM = '"RM"#,##0.00'  # default; spec["currency_format"] overrides
L = {}  # translatable labels; spec["labels"] overrides
def lab(k, d): return L.get(k, d)
MDF = '0.0'

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FILLS = {"normal": None, "flag": AMBER, "free": GREEN}
TEXTC = {"normal": GREYTEXT, "flag": DKAMBER, "free": DKGREEN}


def arial(size=10, bold=False, color="000000", italic=False):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)


def build_workbook(spec, path):
    global RM, L
    RM = spec.get("currency_format", RM)
    L = spec.get("labels", {})
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"

    # Column A carries both the item number and the meta labels ("Prepared By"),
    # so it needs room for the longer of the two or the label clips.
    for col, width in {"A": 15, "B": 62, "C": 9, "D": 14, "E": 16}.items():
        ws.column_dimensions[col].width = width

    # ---- title banner -------------------------------------------------
    ws.merge_cells("A1:E1")
    ws["A1"] = spec.get("title", "THEWEBXITE — BILL OF MATERIALS")
    ws["A1"].font = arial(14, True, "FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    r = 2
    if spec.get("subtitle"):
        ws.merge_cells("A2:E2")
        ws["A2"] = spec["subtitle"]
        ws["A2"].font = arial(11, True, NAVY)
        ws["A2"].alignment = Alignment(horizontal="center")
        ws.row_dimensions[2].height = 20
        r = 3

    # ---- meta block ---------------------------------------------------
    r += 1
    for label, value in spec.get("meta", {}).items():
        ws.cell(r, 1, label).font = arial(10, True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.cell(r, 2, value).font = arial(10)
        ws.cell(r, 2).alignment = Alignment(vertical="center", wrap_text=True)
        r += 1

    # ---- editable rate cell -------------------------------------------
    r += 1
    rate_row = r
    ws.cell(r, 1, lab("rate", "RATE")).font = arial(10, True)
    ws.cell(r, 2, lab("rate_desc", "Man-day rate")).font = arial(10, italic=True)
    cell = ws.cell(r, 4, spec.get("rate", 800))
    # No fill: the BOM is client-facing and a highlighted cell reads as a
    # marked-up draft. The cell is still the single editable input.
    cell.font = arial(11, True, NAVY)
    cell.number_format = RM
    cell.border = BOX
    cell.alignment = Alignment(horizontal="center")
    RATE = f"$D${rate_row}"
    r += 2

    # ---- table header -------------------------------------------------
    header_row = r
    for col, text in [(1, lab("no", "No.")), (2, lab("desc", "Description")), (3, lab("md", "MD")),
                      (4, lab("unit", "Unit Price")), (5, lab("amount", "Amount"))]:
        c = ws.cell(r, col, text)
        c.font = arial(10, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[r].height = 20
    r += 1

    def write_item(row, number, title, md, bullets, kind, free=False):
        ws.cell(row, 1, number).font = arial(10, True)
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 2, title).font = arial(10, True, DKGREEN if free else "000000")
        ws.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=True)
        m = ws.cell(row, 3, md)
        m.font = arial(10, True)
        m.number_format = MDF
        m.alignment = Alignment(horizontal="center", vertical="top")
        u = ws.cell(row, 4, f"={RATE}" if md else "")
        u.font = arial(10)
        u.number_format = RM
        u.alignment = Alignment(horizontal="right", vertical="top")
        a = ws.cell(row, 5, f"=C{row}*{RATE}")
        a.font = arial(10, True, DKGREEN if free else "000000")
        a.number_format = RM
        a.alignment = Alignment(horizontal="right", vertical="top")
        for col in range(1, 6):
            ws.cell(row, col).border = BOX
        row += 1
        for b in bullets:
            btext, bkind = (b, kind) if isinstance(b, str) else (b["text"], b.get("kind", kind))
            marker = "▶  " if bkind == "flag" else ("✓  " if bkind == "free" else "•  ")
            bc = ws.cell(row, 2, marker + btext)
            bc.font = arial(9, bkind != "normal", TEXTC[bkind])
            # No explicit row height: LibreOffice must auto-size wrapped text,
            # otherwise bullets overlap the line item below in the PDF.
            bc.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
            for col in range(1, 6):
                ws.cell(row, col).border = BOX
                if FILLS[bkind]:
                    ws.cell(row, col).fill = PatternFill("solid", fgColor=FILLS[bkind])
            row += 1
        return row

    def subtotal(row, label, first, last, free=False):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        c = ws.cell(row, 1, f"{lab('subtotal', 'Subtotal')} — {label}")
        c.font = arial(10, True)
        c.alignment = Alignment(horizontal="right")
        mds = [f"C{x}" for x in range(first, last + 1)
               if isinstance(ws.cell(x, 3).value, (int, float))]
        amts = [f"E{x}" for x in range(first, last + 1)
                if isinstance(ws.cell(x, 3).value, (int, float))]
        t = ws.cell(row, 3, "=" + "+".join(mds) if mds else 0)
        t.font = arial(10, True)
        t.number_format = MDF
        t.alignment = Alignment(horizontal="center")
        v = ws.cell(row, 5, "=" + "+".join(amts) if amts else 0)
        v.font = arial(10, True)
        v.number_format = RM
        v.alignment = Alignment(horizontal="right")
        fill = GREEN if free else LTBLUE
        for col in range(1, 6):
            ws.cell(row, col).border = BOX
            ws.cell(row, col).fill = PatternFill("solid", fgColor=fill)
        return row + 1

    def phase_header(row, name, fill, color):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        h = ws.cell(row, 1, name)
        h.font = arial(11, True, color)
        h.alignment = Alignment(vertical="center")
        for col in range(1, 6):
            ws.cell(row, col).border = BOX
            ws.cell(row, col).fill = PatternFill("solid", fgColor=fill)
        ws.row_dimensions[row].height = 22
        return row + 1

    def render_sections(row, sections):
        nonlocal item_no
        rows = []
        for sec in sections:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            h = ws.cell(row, 1, sec["name"])
            h.font = arial(10, True, NAVY)
            for col in range(1, 6):
                ws.cell(row, col).border = BOX
                ws.cell(row, col).fill = PatternFill("solid", fgColor=GREY)
            ws.row_dimensions[row].height = 18
            row += 1
            first = row
            for item in sec["items"]:
                item_no += 1
                row = write_item(row, item_no, item["title"], item.get("md", 0),
                                 item.get("bullets", []), item.get("kind", "normal"))
            row = subtotal(row, sec["name"], first, row - 1)
            rows.append(row - 1)
        return row, rows

    # ---- phases -------------------------------------------------------
    phase_subtotal_rows = []
    item_no = 0
    for phase in spec.get("phases", []):
        if phase.get("name"):
            r = phase_header(r, phase["name"], NAVY, "FFFFFF")
        if "sections" in phase:
            r, rows = render_sections(r, phase["sections"])
            phase_subtotal_rows.extend(rows)
        else:
            first = r
            for item in phase["items"]:
                item_no += 1
                r = write_item(r, item_no, item["title"], item.get("md", 0),
                               item.get("bullets", []), item.get("kind", "normal"))
            label = phase["name"].split("\u2014")[-1].strip() if "\u2014" in phase["name"] else phase["name"]
            r = subtotal(r, label, first, r - 1)
            phase_subtotal_rows.append(r - 1)

    # ---- totals -------------------------------------------------------
    r += 1

    def total(label, md_formula, amt_formula, bold=True, fill=LTBLUE, size=10):
        nonlocal r
        # Label spans A:B and the man-day figure sits in C, matching the MD
        # column header and the section subtotal rows above.
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(r, 1, label)
        c.font = arial(size, bold, NAVY if bold else "000000")
        c.alignment = Alignment(horizontal="right", vertical="center")
        if md_formula:
            m = ws.cell(r, 3, md_formula)
            m.font = arial(size, bold)
            m.number_format = '0.0" MD"'
            m.alignment = Alignment(horizontal="center", vertical="center")
        a = ws.cell(r, 5, amt_formula)
        a.font = arial(size, bold, NAVY if bold else "000000")
        a.number_format = RM
        a.alignment = Alignment(horizontal="right", vertical="center")
        for col in range(1, 6):
            ws.cell(r, col).border = BOX
            ws.cell(r, col).fill = PatternFill("solid", fgColor=fill)
        ws.row_dimensions[r].height = 20
        r += 1

    md_sum = "=" + "+".join(f"C{x}" for x in phase_subtotal_rows)
    amt_sum = "=" + "+".join(f"E{x}" for x in phase_subtotal_rows)
    total(lab("total_md", "TOTAL MAN-DAYS / SUBTOTAL"), md_sum, amt_sum)
    sub_row = r - 1

    pct = spec.get("deposit_pct", 50)
    total(lab("deposit", "{pct}% DEPOSIT ON CONFIRMATION").format(pct=pct), None,
          f"=E{sub_row}*{pct / 100}", bold=False, fill="FFFFFF")
    total(lab("balance", "{pct}% BALANCE ON COMPLETION").format(pct=100 - pct), None,
          f"=E{sub_row}*{(100 - pct) / 100}", bold=False, fill="FFFFFF")
    total(lab("total", "TOTAL PROJECT VALUE"), None, f"=E{sub_row}", bold=True, fill=LTBLUE, size=12)

    # ---- notes --------------------------------------------------------
    r += 1
    for note in spec.get("notes", []):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.cell(r, 1, note).font = arial(9, italic=True, color="595959")
        ws.cell(r, 1).alignment = Alignment(vertical="top", wrap_text=True)
        r += 1

    # ---- optional phase, quoted outside the project total -------------
    opt = spec.get("optional_phase")
    if opt:
        r += 2
        # Same neutral treatment as the main phase — the header text and the
        # note carry the "optional" message, not a colour wash.
        r = phase_header(r, opt["name"], NAVY, "FFFFFF")
        secs = opt.get("sections") or [{"name": opt["name"], "items": opt["items"]}]
        r, opt_rows = render_sections(r, secs)

        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        n = ws.cell(r, 1, lab("opt_note", "Not included in the total project value above. Quoted separately "
                          "and may be commissioned at any time."))
        n.font = arial(9, italic=True, color="595959")
        n.alignment = Alignment(vertical="top", wrap_text=True)
        r += 2

        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(r, 1, lab("opt_total", "TOTAL IF PHASE 2 COMMISSIONED"))
        c.font = arial(10, True, NAVY)
        c.alignment = Alignment(horizontal="right", vertical="center")
        m = ws.cell(r, 3, "=C{0}+{1}".format(sub_row, "+".join(f"C{x}" for x in opt_rows)))
        m.font = arial(10, True)
        m.number_format = '0.0" MD"'
        m.alignment = Alignment(horizontal="center", vertical="center")
        a = ws.cell(r, 5, "=E{0}+{1}".format(sub_row, "+".join(f"E{x}" for x in opt_rows)))
        a.font = arial(10, True, NAVY)
        a.number_format = RM
        a.alignment = Alignment(horizontal="right", vertical="center")
        for col in range(1, 6):
            ws.cell(r, col).border = BOX
            ws.cell(r, col).fill = PatternFill("solid", fgColor=LTBLUE)
        ws.row_dimensions[r].height = 20
        r += 1

    last_row = r - 1

    # ---- print setup --------------------------------------------------
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.45
    ws.page_margins.right = 0.45
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.6
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.3
    ws.print_area = f"A1:E{last_row}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.print_options.horizontalCentered = True
    ws.print_options.gridLines = False

    footer_name = spec.get("output_name", "Bill of Materials")
    ws.oddFooter.left.text = f"TheWebxite — {footer_name}"
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.color = "808080"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.color = "808080"

    wb.save(path)
    return last_row


def clear_row_heights(path):
    """Release any fixed row heights so wrapped bullets don't overlap in PDF."""
    wb = load_workbook(path)
    ws = wb["BOM"]
    for row in list(ws.row_dimensions):
        rd = ws.row_dimensions[row]
        if rd.height is not None and row > 1:
            # keep the banner rows, free everything that holds wrapped text
            if ws.cell(row, 2).alignment.wrap_text:
                rd.height = None
    wb.save(path)


def to_pdf(xlsx_path, outdir, output_name):
    helper = "/mnt/skills/public/xlsx/scripts/office/soffice.py"
    tmp = tempfile.mkdtemp()
    if os.path.exists(helper):
        cmd = [sys.executable, helper, "--convert-to", "pdf", "--outdir", tmp, xlsx_path]
    else:
        cmd = ["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, xlsx_path]
    subprocess.run(cmd, check=True, capture_output=True)
    produced = os.path.join(tmp, os.path.splitext(os.path.basename(xlsx_path))[0] + ".pdf")
    target = os.path.join(outdir, f"{output_name}.pdf")
    shutil.copy(produced, target)
    shutil.rmtree(tmp, ignore_errors=True)
    return target


def main():
    ap = argparse.ArgumentParser(description="Build a TWX BOM as xlsx + pdf")
    ap.add_argument("--spec", required=True, help="path to the BOM spec JSON")
    ap.add_argument("--outdir", default="/mnt/user-data/outputs")
    args = ap.parse_args()

    with open(args.spec) as f:
        spec = json.load(f)

    os.makedirs(args.outdir, exist_ok=True)
    name = spec.get("output_name", "Bill of Materials")
    xlsx = os.path.join(args.outdir, f"{name}.xlsx")

    build_workbook(spec, xlsx)
    clear_row_heights(xlsx)
    pdf = to_pdf(xlsx, args.outdir, name)

    print(f"xlsx: {xlsx}")
    print(f"pdf:  {pdf}")
    print("Verify the PDF visually before presenting — check bullets don't "
          "overlap the line item beneath them.")


if __name__ == "__main__":
    main()
